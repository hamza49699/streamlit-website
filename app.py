import streamlit as st  
from pdf2docx import Converter 
from docx import Document
from fpdf import FPDF 
import io
import time

st.set_page_config(page_title="📄 PDF ↔ Word Converter", page_icon="🔄")

st.title("📄 Free Professional PDF ↔ Word Converter 🔄")
st.write("📢 Convert between PDF and Word files quickly and easily! 🎯")

# Select conversion type
conversion_type = st.radio("Select conversion type:", ["PDF to Word", "Word to PDF"])

uploaded_file = st.file_uploader("📤 Upload your file", type=["pdf", "docx"] if conversion_type == "Word to PDF" else ["pdf"])

if uploaded_file:
    try:
        if conversion_type == "PDF to Word":
            if "converted_file" not in st.session_state or "last_uploaded_file" not in st.session_state or st.session_state["last_uploaded_file"] != uploaded_file.name:
                with open("temp.pdf", "wb") as f:
                    f.write(uploaded_file.read())

                with st.spinner("⏳ Converting your PDF to Word... Please wait. ⚡"):
                    time.sleep(2)
                    
                    word_io = io.BytesIO()
                    cv = Converter("temp.pdf")
                    cv.convert(word_io, start=0, end=None)
                    cv.close()
                    word_io.seek(0)

                    st.session_state["converted_file"] = word_io
                    st.session_state["last_uploaded_file"] = uploaded_file.name
            
            st.success("🎉✅ Conversion Successful! Download your Word file below. 📂")
            st.download_button(
                label="📥 Download Word Document",
                data=st.session_state["converted_file"].getvalue(),
                file_name="converted_document.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        
        elif conversion_type == "Word to PDF":
            if uploaded_file.name.endswith(".docx"):
                with st.spinner("⏳ Converting your Word document to PDF... Please wait. ⚡"):
                    time.sleep(2)
                    
                    doc = Document(uploaded_file)
                    pdf_io = io.BytesIO()
                    pdf = FPDF()
                    pdf.set_auto_page_break(auto=True, margin=15)
                    pdf.add_page()
                    pdf.set_font("Arial", size=12)
                    
                    for para in doc.paragraphs:
                        pdf.cell(200, 10, txt=para.text, ln=True, align='L')
                    
                    pdf.output(pdf_io, 'F')
                    pdf_io.seek(0)

                    st.success("🎉✅ Conversion Successful! Download your PDF file below. 📂")
                    st.download_button(
                        label="📥 Download PDF Document",
                        data=pdf_io.getvalue(),
                        file_name="converted_document.pdf",
                        mime="application/pdf",
                    )
            else:
                st.error("❌⚠️ Please upload a valid .docx file for conversion.")
    
    except Exception as e:
        st.error(f"❌⚠️ Error: {str(e)}")

st.markdown("---")
st.markdown(
    "🔧 Developed by **Muhammad Hamza Javed** | 💜 Follow me on [GitHub](https://github.com/hamza49699) & [LinkedIn](https://www.linkedin.com/in/hamza-khan-7472b822b/) "
)
st.markdown("© 2025 PDF ↔ Word Converter. All rights reserved.")
